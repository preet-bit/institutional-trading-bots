"""
Add new strategy results (P06-P15) to the EURUSD Excel workbook.
Appends 3 sheets: New_Strategies_Leaderboard, New_Walkforward_Windows, New_Yearly_Returns
"""
import os, sys
import numpy as np, pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from eurusd_new_strategies import run_new_strategies


def write_sheet(ws, d, fmts=None, freeze='C2'):
    fmts = fmts or {}
    FONT = 'Arial'
    hdr_fill = PatternFill('solid', fgColor='1F3864')
    hdr_font = Font(name=FONT, bold=True, color='FFFFFF', size=10)
    base = Font(name=FONT, size=10)
    thin = Side(style='thin', color='D9D9D9')
    border = Border(bottom=thin)

    for j, col in enumerate(d.columns, 1):
        cell = ws.cell(1, j, str(col))
        cell.font = hdr_font; cell.fill = hdr_fill
        cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    ws.row_dimensions[1].height = 42

    for i, row in enumerate(d.itertuples(index=False), 2):
        for j, v in enumerate(row, 1):
            if isinstance(v, (np.floating, float)):
                v = None if (np.isnan(v) or np.isinf(v)) else float(v)
            elif isinstance(v, np.integer): v = int(v)
            elif isinstance(v, (np.bool_, bool)): v = 'Y' if v else 'N'
            cell = ws.cell(i, j, v)
            cell.font = base; cell.border = border
            col_name = d.columns[j - 1]
            if col_name in fmts and v is not None:
                cell.number_format = fmts[col_name]

    for j, col in enumerate(d.columns, 1):
        mx = max([len(str(col)) * 0.6] + [len(str(x)) for x in d[col].head(60)])
        ws.column_dimensions[get_column_letter(j)].width = min(max(9, mx + 2), 48)
    ws.freeze_panes = freeze


def main():
    EV, WF, YR, res, OOS = run_new_strategies()

    xlsx_path = os.path.abspath(os.path.join(
        os.path.dirname(__file__), '..', 'EURUSD_4h_50_strategies_backtest.xlsx'))
    wb = openpyxl.load_workbook(xlsx_path)

    # Remove old sheets if present
    for sname in ['New_Strategies_Leaderboard', 'New_Walkforward_Windows', 'New_Yearly_Returns']:
        if sname in wb.sheetnames:
            del wb[sname]

    fm_std = {
        'win_rate': '0.0%', 'payoff': '0.00', 'profit_factor': '0.00',
        'sharpe': '0.00', 'sortino': '0.00', 'total_ret_pct': '0.0',
        'ann_ret_pct': '0.0', 'max_dd_pct': '0.0', 'dsr_prob': '0.000',
        'exp_max_sr': '0.000',
        'pct_windows_profitable': '0.0%', 'unleveraged_1x_ret_pct': '0.0',
        'leverage_3x_ret_pct': '0.0', 'compounded_1_5pct_risk_pct': '0.0',
        'compounded_2_0pct_risk_pct': '0.0'
    }

    ws_lb = wb.create_sheet('New_Strategies_Leaderboard')
    lb_cols = [
        'strategy_id', 'strategy', 'family', 'trades', 'win_rate', 'payoff',
        'profit_factor', 'sharpe', 'sortino', 'max_dd_pct', 'ann_ret_pct',
        'windows_traded', 'windows_profitable', 'pct_windows_profitable',
        'unleveraged_1x_ret_pct', 'leverage_3x_ret_pct',
        'compounded_1_5pct_risk_pct', 'compounded_2_0pct_risk_pct',
        'dsr_prob', 'exp_max_sr', 'verdict'
    ]
    avail = [c for c in lb_cols if c in EV.columns]
    write_sheet(ws_lb, EV[avail], fmts=fm_std)

    ws_wf = wb.create_sheet('New_Walkforward_Windows')
    wf_cols = [
        'strategy_id', 'strategy', 'family', 'window', 'train', 'test',
        'selected', 'is_breadth_profitable', 'is_trades', 'is_win_rate',
        'is_sharpe', 'is_pf', 'is_ret_pct', 'oos_trades', 'oos_win_rate',
        'oos_pf', 'oos_ret_pct', 'oos_exp_R', 'oos_sharpe'
    ]
    wf_avail = [c for c in wf_cols if c in WF.columns]
    write_sheet(ws_wf, WF[wf_avail], fmts={
        'is_breadth_profitable': '0.0%', 'is_win_rate': '0.0%', 'oos_win_rate': '0.0%',
        'is_sharpe': '0.00', 'is_pf': '0.00', 'is_ret_pct': '0.0',
        'oos_pf': '0.00', 'oos_ret_pct': '0.0', 'oos_exp_R': '0.000', 'oos_sharpe': '0.00'
    })

    ws_yr = wb.create_sheet('New_Yearly_Returns')
    write_sheet(ws_yr, YR, fmts={str(y): '0.0;[Red]-0.0' for y in range(2013, 2027)})

    wb.save(xlsx_path)
    print(f"\nSaved to {xlsx_path}")

    out_path = os.path.abspath(os.path.join(
        os.path.dirname(__file__), '..', 'outputs', 'EURUSD_4h_50_strategies_backtest.xlsx'))
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)
    print(f"Copied to {out_path}")


if __name__ == '__main__':
    main()
