import pickle, numpy as np, pandas as pd, shutil, os, zipfile
import matplotlib; matplotlib.use('Agg')
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from core import COST_RT, FEE, SLIP
from strategies import STRATS, SL_GRID, RR_GRID

A = pickle.load(open('analysis.pkl', 'rb')); RES = pickle.load(open('results.pkl', 'rb'))['res']
df = pd.read_pickle('df.pkl'); T = pd.to_datetime(df.t.values)
EV, WF, CS, YR, PORT, CUR, DEFS, CD = [A[k] for k in ('EV', 'WF', 'CS', 'YR', 'PORT', 'CUR', 'DEFS', 'CD')]
OUT = os.environ.get('OUTPUT_DIR', os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'outputs')))
os.makedirs(OUT, exist_ok=True)

EV = EV.copy()
nt0 = EV.trades == 0
EV.loc[nt0, 'verdict'] = 'FAIL (no in-sample edge found, never traded)'
EV['verdict'] = EV.verdict.str.replace('WATCH (profitable, not statistically robust)', 'WATCH (profitable OOS, not statistically robust)', regex=False)
lb = EV.sort_values(['sharpe', 'trades'], ascending=[False, False]).reset_index(drop=True)

# ------------------------------------------------------------------ OOS trade log
rows = []
for wi, w in WF.iterrows():
    if not isinstance(w.get('sig_params'), str): continue
    sid = w.strategy_id; key = (sid, int(w.param_idx), w.sl_atr, w.rr)
    ei, xi, dr, g, rk = RES[key]
    a, b = w.test.split('..'); a = pd.Timestamp(a); b = pd.Timestamp(b) + pd.Timedelta(days=1)
    et = T[ei]; m = (et >= a) & (et < b)
    for i in np.where(m)[0]:
        rows.append((sid, int(w.window), w.selected, T[ei[i]], T[xi[i]], 'LONG' if dr[i] == 1 else 'SHORT', df.o.values[ei[i]],
                     g[i] * 100, (g[i] - COST_RT) * 100, rk[i] * 100, (g[i] - COST_RT) / rk[i]))
TL = pd.DataFrame(rows, columns=['strategy_id', 'wf_window', 'params_used', 'entry_time', 'exit_time', 'side', 'entry_price',
                                 'gross_ret_pct', 'net_ret_pct', 'stop_distance_pct', 'net_R'])
TL.to_csv(f'{OUT}/oos_trade_log_all_strategies.csv', index=False)

# ------------------------------------------------------------------ win rate vs RR (full-sample, all combos, unselected)
wr = []
for sl in SL_GRID:
    for rr in RR_GRID:
        n = 0; wg = 0; wn = 0; sg = 0.0; sn = 0.0
        for st in STRATS:
            for pi in range(len(st['grid'])):
                ei, xi, dr, g, rk = RES[(st['id'], pi, sl, rr)]
                n += len(g); wg += (g > 0).sum(); wn += ((g - COST_RT) > 0).sum(); sg += (g / rk).sum(); sn += ((g - COST_RT) / rk).sum()
        wr.append(dict(sl_atr=sl, rr_target=rr, trades=n, win_rate_gross=wg / n, win_rate_net=wn / n, theoretical_breakeven_wr=1 / (1 + rr),
                       expectancy_R_gross=sg / n, expectancy_R_net=sn / n))
WRR = pd.DataFrame(wr)

# ------------------------------------------------------------------ charts
d0o, d1o = A['d0o'], A['d1o']; idx = pd.date_range('2018-01-01', periods=d1o)[d0o:d1o]
plt.rcParams.update({'font.family': 'DejaVu Sans', 'axes.grid': True, 'grid.alpha': .25})
fig, ax = plt.subplots(figsize=(11, 5.5))
ax.plot(idx, np.cumsum(A['allD']) * 100, label='Equal-weight all 50 strategies (OOS, net of costs)', lw=2, color='#c0392b')
ax.plot(idx, np.cumsum(A['bh']) * 100, label='BTC buy & hold (arithmetic cum. return)', lw=1.5, color='#7f8c8d')
ax.axhline(0, color='k', lw=.8); ax.set_ylabel('Cumulative return, % of notional'); ax.set_title('Walk-forward out-of-sample 2020-01 to 2026-06: all 50 strategies vs buy & hold')
ax.legend(); fig.tight_layout(); fig.savefig(f'{OUT}/chart1_portfolio_vs_buyhold.png', dpi=140); plt.close(fig)

top = lb[lb.trades > 0].head(8)
fig, ax = plt.subplots(figsize=(11, 5.5))
for _, r in top.iterrows(): ax.plot(idx, np.cumsum(A['DAILY'][r.strategy_id]) * 100, lw=1.3, label=f"{r.strategy_id} {r.strategy} (Sharpe {r.sharpe:.2f})")
ax.axhline(0, color='k', lw=.8); ax.set_ylabel('Cumulative net return, % of notional'); ax.set_title('Top strategies by walk-forward OOS Sharpe (selected AFTER the fact, so flattering)')
ax.legend(fontsize=7.5); fig.tight_layout(); fig.savefig(f'{OUT}/chart2_top_strategies_oos_equity.png', dpi=140); plt.close(fig)

t2 = EV[EV.trades > 0]
fig, ax = plt.subplots(figsize=(7.5, 6.5))
ax.scatter(t2.mean_is_sharpe, t2.sharpe, s=45, color='#2c3e50')
for _, r in t2.iterrows(): ax.annotate(r.strategy_id, (r.mean_is_sharpe, r.sharpe), fontsize=7, xytext=(3, 3), textcoords='offset points')
lim = [min(-1, t2.sharpe.min() - .1), max(t2.mean_is_sharpe.max(), 1) + .1]; ax.plot(lim, lim, 'r--', lw=1, label='no overfit (OOS = in-sample)')
ax.axhline(0, color='k', lw=.8); ax.set_xlabel('Mean in-sample Sharpe of the parameters chosen (training windows)'); ax.set_ylabel('Out-of-sample Sharpe (test windows)')
ax.set_title('Overfitting gap: in-sample vs out-of-sample Sharpe'); ax.legend(); fig.tight_layout(); fig.savefig(f'{OUT}/chart3_in_sample_vs_oos_sharpe.png', dpi=140); plt.close(fig)

# ------------------------------------------------------------------ workbook
FONT = 'Arial'
hdr_fill = PatternFill('solid', fgColor='1F3864'); hdr_font = Font(name=FONT, bold=True, color='FFFFFF', size=10)
base = Font(name=FONT, size=10); bold = Font(name=FONT, size=10, bold=True)
thin = Side(style='thin', color='D9D9D9'); border = Border(bottom=thin)
wb = Workbook()

def write_df(ws, d, r0=1, fmts=None, widths=None, freeze=True):
    fmts = fmts or {}
    for j, c in enumerate(d.columns, 1):
        cell = ws.cell(r0, j, c); cell.font = hdr_font; cell.fill = hdr_fill
        cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    ws.row_dimensions[r0].height = 42
    for i, row in enumerate(d.itertuples(index=False), r0 + 1):
        for j, v in enumerate(row, 1):
            if isinstance(v, (np.floating, float)):
                v = None if (np.isnan(v) or np.isinf(v)) else float(v)
            elif isinstance(v, (np.integer,)): v = int(v)
            elif isinstance(v, (np.bool_, bool)): v = 'Y' if v else 'N'
            cell = ws.cell(i, j, v); cell.font = base; cell.border = border
            f = fmts.get(d.columns[j - 1])
            if f and v is not None: cell.number_format = f
    for j, c in enumerate(d.columns, 1):
        w = (widths or {}).get(c)
        if w is None:
            mx = max([len(str(c)) * .6] + [len(str(x)) for x in d[c].head(60)]); w = min(max(9, mx + 2), 44)
        ws.column_dimensions[get_column_letter(j)].width = w
    if freeze: ws.freeze_panes = ws.cell(r0 + 1, 3)
    return ws

# ---- Leaderboard first (Summary refers to it)
ws_s = wb.active; ws_s.title = 'Summary'
ws = wb.create_sheet('Leaderboard')
cols = [('Rank', None), ('ID', 'strategy_id'), ('Strategy', 'strategy'), ('Family', 'family'), ('OOS trades', 'trades'), ('Win rate', 'win_rate'),
        ('Avg win (R)', 'avg_win_R'), ('Avg loss (R)', 'avg_loss_R'), ('Realised payoff (avg win / avg loss)', 'payoff'), ('Breakeven win rate', 'F'),
        ('Edge: win rate - breakeven', 'F2'), ('Expectancy (R / trade)', 'expectancy_R'), ('Profit factor', 'profit_factor'), ('Sharpe (ann.)', 'sharpe'),
        ('Sortino (ann.)', 'sortino'), ('Ann. return (% of notional)', 'ann_ret_pct'), ('Max drawdown (% of notional)', 'max_dd_pct'), ('Calmar', 'calmar'),
        ('% test windows profitable', 'pct_windows_profitable'), ('Deflated Sharpe prob.', 'dsr'), ('Bootstrap p (mean<=0)', 'p_value'),
        ('Holm-adj. p (x50)', 'p_holm_adj'), ('Gates passed (of 4)', 'gates_passed'), ('Verdict', 'verdict')]
for j, (h, _) in enumerate(cols, 1):
    c = ws.cell(1, j, h); c.font = hdr_font; c.fill = hdr_fill; c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
ws.row_dimensions[1].height = 54
fm = {'win_rate': '0.0%', 'avg_win_R': '0.00', 'avg_loss_R': '0.00', 'payoff': '0.00', 'expectancy_R': '0.000', 'profit_factor': '0.00', 'sharpe': '0.00',
      'sortino': '0.00', 'ann_ret_pct': '0.0', 'max_dd_pct': '0.0', 'calmar': '0.00', 'pct_windows_profitable': '0%', 'dsr': '0.000', 'p_value': '0.000', 'p_holm_adj': '0.000'}
for i, r in lb.iterrows():
    rr_ = i + 2
    for j, (h, k) in enumerate(cols, 1):
        if h == 'Rank': v = i + 1
        elif k == 'F': v = f'=IF(ISNUMBER(I{rr_}),1/(1+I{rr_}),"")'
        elif k == 'F2': v = f'=IF(AND(ISNUMBER(F{rr_}),ISNUMBER(J{rr_})),F{rr_}-J{rr_},"")'
        else:
            v = r[k]
            if isinstance(v, (float, np.floating)): v = None if (np.isnan(v) or np.isinf(v)) else float(v)
            elif isinstance(v, np.integer): v = int(v)
        c = ws.cell(rr_, j, v); c.font = base; c.border = border
        if k in fm: c.number_format = fm[k]
        if k == 'F': c.number_format = '0.0%'
        if k == 'F2': c.number_format = '+0.0%;-0.0%;0.0%'
for j, w in enumerate([6, 7, 32, 11, 9, 9, 9, 9, 13, 11, 12, 12, 9, 9, 9, 12, 12, 8, 12, 11, 11, 10, 9, 44], 1): ws.column_dimensions[get_column_letter(j)].width = w
ws.freeze_panes = 'D2'; ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{len(lb) + 1}'
ws.conditional_formatting.add(f'X2:X{len(lb) + 1}', FormulaRule(formula=['LEFT(X2,4)="PASS"'], fill=PatternFill('solid', bgColor='C6EFCE')))
ws.conditional_formatting.add(f'X2:X{len(lb) + 1}', FormulaRule(formula=['LEFT(X2,5)="WATCH"'], fill=PatternFill('solid', bgColor='FFEB9C')))
ws.conditional_formatting.add(f'X2:X{len(lb) + 1}', FormulaRule(formula=['LEFT(X2,4)="FAIL"'], fill=PatternFill('solid', bgColor='F4CCCC')))
LBN = len(lb) + 1
n_ = len(lb) + 3
notes = ['Sharpe etc. are walk-forward OUT-OF-SAMPLE: every trade was taken with parameters chosen only from the prior 24 months. Net of 0.12% round-trip costs.',
         'Strategies with 0 trades found no parameter set that was profitable in-sample in ANY training window, so the walk-forward rule kept them flat. That is a real result, not a bug.',
         'Gates (fixed in advance): >=100 OOS trades; profit factor >=1.10 AND Sharpe >=0.5; >=60% of traded test windows profitable; Deflated Sharpe probability >=0.95. PASS needs all 4.',
         'Breakeven win rate = 1/(1+realised payoff). Edge = actual win rate minus that. Columns J and K are live formulas.']
for k, t in enumerate(notes): c = ws.cell(n_ + k, 3, t); c.font = Font(name=FONT, size=9, italic=True)

# ---- Summary
S = ws_s; S.sheet_view.showGridLines = False
S.column_dimensions['A'].width = 3; S.column_dimensions['B'].width = 46; S.column_dimensions['C'].width = 18; S.column_dimensions['D'].width = 90
def put(r, b=None, c=None, d=None, f=None, fmt=None):
    for col, v in ((2, b), (3, c), (4, d)):
        if v is not None:
            cell = S.cell(r, col, v); cell.font = f or base; cell.alignment = Alignment(wrap_text=True, vertical='top')
            if col == 3 and fmt: cell.number_format = fmt
r = 2; S.cell(r, 2, 'BTC/USD 15-minute backtest: 50 strategies, walk-forward validated').font = Font(name=FONT, size=15, bold=True)
r = 4; put(r, 'HEADLINE RESULT', f=bold)
put(5, 'Strategies tested', f'=COUNTA(Leaderboard!B2:B{LBN})')
put(6, 'PASS (all 4 robustness gates)', f'=COUNTIF(Leaderboard!X2:X{LBN},"PASS*")', 'Profitable, consistent and statistically robust after correcting for testing 50 strategies')
put(7, 'WATCH (profitable OOS, not robust)', f'=COUNTIF(Leaderboard!X2:X{LBN},"WATCH*")', 'Made money out-of-sample but cannot be distinguished from luck')
put(8, 'FAIL', f'=COUNTIF(Leaderboard!X2:X{LBN},"FAIL*")', 'Lost money net of costs, or never found an in-sample edge')
put(9, 'Strategies that never traded OOS', f'=COUNTIF(Leaderboard!E2:E{LBN},0)', 'No parameter set was profitable in training in any window')
put(10, 'Strategies with positive OOS Sharpe', f'=COUNTIF(Leaderboard!N2:N{LBN},">0")')
put(11, 'Best OOS Sharpe (annualised)', f'=MAX(Leaderboard!N2:N{LBN})', None, fmt='0.00')
put(12, 'Median OOS Sharpe of the 50', f'=MEDIAN(Leaderboard!N2:N{LBN})', None, fmt='0.00')
put(13, 'Equal-weight all-50 OOS Sharpe', float(PORT.sharpe.iloc[0]), 'No selection bias. Hard-coded from analysis; see Portfolios sheet.', fmt='0.00')
put(14, 'BTC buy & hold Sharpe, same period', float(PORT.sharpe.iloc[1]), 'The benchmark every strategy here failed to beat', fmt='0.00')
S.cell(16, 2, 'WHAT THIS MEANS').font = bold
best = lb[lb.trades > 0].iloc[0]
_w = WF[(WF.strategy_id == best.strategy_id) & (WF.oos_trades.fillna(0) > 0)].oos_ret_pct.values[-4:]
late = f"In its last 4 traded windows it lost money in {int((_w < 0).sum())}, and no parameter set qualified in the final {int((WF[WF.strategy_id == best.strategy_id].selected.str.startswith('NO TRADE')).tail(2).sum())} windows: the edge is fading. Treat it as unproven."
_g = WRR.groupby('rr_target').agg(w=('win_rate_net', 'mean'), e=('expectancy_R_net', 'mean'))
wr_txt = f"RR 1 -> {_g.w[1.0]:.0%} win rate, RR 2 -> {_g.w[2.0]:.0%}, RR 3 -> {_g.w[3.0]:.0%}"
txt = [
 f"None of the 50 strategies cleared all four robustness gates. The best was {best.strategy_id} {best.strategy}: OOS Sharpe {best.sharpe:.2f}, win rate {best.win_rate:.1%}, realised payoff {best.payoff:.2f}, profit factor {best.profit_factor:.2f}, max drawdown {best.max_dd_pct:.0f}% of notional, deflated-Sharpe probability {best.dsr:.2f}. {late}",
 "Costs are the main killer. Average gross edge across all unselected parameter sets is about 0 bps per trade against a 12 bps round-trip cost. On average only about 2% of a strategy's parameter sets are net-positive over the full sample (see Cost_Diagnostics). With half the cost, a few strategies become interesting; with zero cost, several look good. That is exactly why 1-minute / 15-minute scalping is hard.",
 "No strategy combined a high win rate, a high payoff AND a robust edge out-of-sample (the best has a ~47% win rate with a ~1.4 payoff). See WinRate_vs_RR: raising the target multiple lowers the hit rate almost exactly as theory predicts (pooled net: " + wr_txt + "), and net expectancy is negative at every setting.",
 "Optimising and then testing on the same data would have shown much better numbers. OOS_Full_Metrics (mean_is_sharpe, wfe, fs_* columns) and chart 3 show in-sample Sharpe well above OOS Sharpe for the strategies that did trade. I deliberately did NOT keep tuning until something passed, because that would produce a result that looks great and fails live.",
]
for k, t in enumerate(txt):
    S.cell(17 + k, 2, t).font = base; S.merge_cells(start_row=17 + k, start_column=2, end_row=17 + k, end_column=4)
    S.cell(17 + k, 2).alignment = Alignment(wrap_text=True, vertical='top'); S.row_dimensions[17 + k].height = 62
S.cell(22, 2, 'DATA').font = bold
data_rows = [('File', 'btc_15m_data_2018_to_2025.csv', 'Despite the name the file runs 2018-01-01 00:00 to 2026-06-15 23:15 UTC'),
             ('Bars used', 295914, '15-minute OHLCV + taker-buy volume (Binance-style BTCUSDT klines)'),
             ('Cleaning', '2 rows removed', 'Two corrupt rows on 2018-07-07 06:00 and 06:15 with prices around 108,000 (about 16x the real price) duplicated timestamps. All other large moves (e.g. 2018-01-11, 2020-03-12/13) were checked and are real.'),
             ('Gaps', '30 small gaps', 'Missing 15m bars (mostly exchange maintenance, up to ~1.4 days once). Not forward-filled. Indicators run over the bars that exist.')]
for k, (a, b, c) in enumerate(data_rows): put(23 + k, a, b, c)
S.cell(28, 2, 'METHOD AND ASSUMPTIONS').font = bold
meth = [('Strategies', '50 across 6 families', 'Trend 13, Breakout 12, Mean reversion 12, Momentum 6, Volume/order-flow 5, Volatility/time 2. Definitions in Strategy_Definitions.'),
        ('Parameter grid per strategy', '4 signal sets x 3 SL x 3 RR', 'Stop = 1, 2 or 3 x ATR(14). Target = 1, 2 or 3 x stop (so RR 1, 2, 3). Max hold 96 bars (24h). Small grids on purpose, to limit data-mining.'),
        ('Execution', 'Next-bar open', 'Signal on bar close, fill at the NEXT bar open. No lookahead. If stop and target are both touched in one bar, the stop is assumed hit first. Gaps through a level fill at the open.'),
        ('Costs', f'{COST_RT * 100:.2f}% round trip', f'Taker fee {FEE * 100:.2f}% + slippage {SLIP * 100:.2f}% per side, deducted from every trade. Funding rates are NOT modelled (24h max hold limits the effect).'),
        ('Position / sizing', '1x notional, one trade at a time', 'No leverage, no compounding, long and short. Returns are % of notional; R = net return / stop distance. Sharpe is on daily P&L (calendar days, including flat days) x sqrt(365).'),
        ('Walk-forward', 'Train 24m / test 6m / roll 6m', f'13 test windows from 2020-01-01 to 2026-06-15. In each window the best parameter set by in-sample Sharpe (needs >=80 trades, PF>1.05) is applied to the unseen next 6 months. Trades that would exit after the train end are purged. If nothing qualifies, the strategy stays flat.'),
        ('Multiple-testing correction', 'Deflated Sharpe, Holm', 'Deflated Sharpe (Bailey & Lopez de Prado) with 50 trials; bootstrap p-values (2,000 resamples of trades) with a x50 Holm/Bonferroni adjustment; Monte-Carlo trade-shuffle drawdowns (500 shuffles).'),
        ('Known limitations', '', 'Single asset, single exchange history. Win/loss on the same bar is approximated. Walk-forward windows reuse the same history I could see, and the 4 gates were fixed before results. Past results do not imply future results. This is research, not financial advice.')]
for k, (a, b, c) in enumerate(meth):
    put(29 + k, a, b, c); S.row_dimensions[29 + k].height = 44
S.cell(38, 2, 'SHEETS').font = bold
for k, (a, b) in enumerate([('Leaderboard', 'All 50 strategies ranked by OOS Sharpe, with verdicts'), ('Strategy_Definitions', 'Logic and parameter grid of each strategy'),
                            ('OOS_Full_Metrics', 'Every metric, including Monte-Carlo percentiles and full-sample comparison'), ('Walkforward_Windows', 'Per strategy per window: chosen params, in-sample vs out-of-sample'),
                            ('Yearly_Returns', 'OOS net return by calendar year'), ('Cost_Sensitivity', 'OOS Sharpe/return at 0x, 0.5x, 1x, 2x costs'), ('Cost_Diagnostics', 'Gross vs net edge across ALL parameter sets'),
                            ('WinRate_vs_RR', 'Why high win rate and high RR do not coexist'), ('Current_Params', 'Parameters fitted on the latest 24 months (not validated)'), ('Portfolios', 'Equal-weight blends vs buy & hold')]):
    put(39 + k, a, None, b)

# ---- other sheets
write_df(wb.create_sheet('Strategy_Definitions'), DEFS, widths={'name': 34, 'logic': 60, 'param_grid': 52})
oc = ['strategy_id', 'strategy', 'family', 'trades', 'win_rate', 'avg_win_R', 'avg_loss_R', 'payoff', 'breakeven_wr', 'edge_vs_breakeven_wr', 'expectancy_R', 'profit_factor', 'total_ret_pct',
      'ann_ret_pct', 'sharpe', 'sortino', 'max_dd_pct', 'calmar', 'windows_traded', 'windows_profitable', 'windows_no_trade', 'pct_windows_profitable', 'mean_is_sharpe', 'wfe', 'mean_is_breadth',
      'param_modal_share', 'dsr', 'sr0_ann', 'p_value', 'p_holm_adj', 'mc_ret_p5_pct', 'mc_ret_p50_pct', 'mc_ret_p95_pct', 'mc_maxdd_median_pct', 'mc_maxdd_p95_pct',
      'fullsample_best', 'fs_trades', 'fs_win_rate', 'fs_pf', 'fs_sharpe', 'fs_ret_pct', 'gates_passed', 'verdict']
pct = {c: '0.0%' for c in ['win_rate', 'breakeven_wr', 'edge_vs_breakeven_wr', 'pct_windows_profitable', 'mean_is_breadth', 'param_modal_share', 'fs_win_rate']}
num2 = {c: '0.00' for c in ['avg_win_R', 'avg_loss_R', 'payoff', 'profit_factor', 'total_ret_pct', 'ann_ret_pct', 'sharpe', 'sortino', 'max_dd_pct', 'calmar', 'mean_is_sharpe', 'wfe', 'sr0_ann',
                            'mc_ret_p5_pct', 'mc_ret_p50_pct', 'mc_ret_p95_pct', 'mc_maxdd_median_pct', 'mc_maxdd_p95_pct', 'fs_pf', 'fs_sharpe', 'fs_ret_pct']}
num3 = {c: '0.000' for c in ['expectancy_R', 'dsr', 'p_value', 'p_holm_adj']}
write_df(wb.create_sheet('OOS_Full_Metrics'), lb[oc], fmts={**pct, **num2, **num3}, widths={'strategy': 32, 'fullsample_best': 30, 'verdict': 40})
wfc = ['strategy_id', 'strategy', 'window', 'train', 'test', 'selected', 'is_breadth_profitable', 'is_trades', 'is_win_rate', 'is_sharpe', 'is_pf', 'is_ret_pct', 'oos_trades', 'oos_win_rate', 'oos_pf', 'oos_ret_pct', 'oos_exp_R', 'oos_sharpe']
write_df(wb.create_sheet('Walkforward_Windows'), WF[wfc], fmts={'is_breadth_profitable': '0%', 'is_win_rate': '0.0%', 'oos_win_rate': '0.0%', 'is_sharpe': '0.00', 'is_pf': '0.00', 'is_ret_pct': '0.0',
          'oos_pf': '0.00', 'oos_ret_pct': '0.0', 'oos_exp_R': '0.000', 'oos_sharpe': '0.00'}, widths={'strategy': 30, 'train': 24, 'test': 24, 'selected': 42})
ws = wb.create_sheet('Yearly_Returns'); write_df(ws, YR, fmts={str(y): '0.0;[Red]-0.0' for y in range(2020, 2027)}, widths={'strategy': 32})
ws.cell(len(YR) + 3, 2, 'OOS net return by exit year, in % of notional (sum of trade returns, no compounding). 2026 is partial (to 15 June).').font = Font(name=FONT, size=9, italic=True)
write_df(wb.create_sheet('Cost_Sensitivity'), CS, fmts={c: '0.00' for c in CS.columns if c.startswith(('sharpe', 'ret_'))}, widths={'strategy': 32})
write_df(wb.create_sheet('Cost_Diagnostics'), CD, fmts={'avg_gross_bps_per_trade': '0.00', 'avg_net_bps_per_trade': '0.00', 'cost_bps_rt': '0.0', 'median_stop_pct': '0.00',
         'pct_combos_gross_positive': '0%', 'pct_combos_net_positive': '0%'}, widths={'strategy': 32})
ws = wb['Cost_Diagnostics']; k = len(CD) + 3
for t in ['Unselected, full-sample 2018-2026, across ALL parameter sets with >=200 trades. 1 bp = 0.01%.', 'Gross = before costs. Net = after the 12 bps round-trip cost. Shows where the edge goes.']:
    ws.cell(k, 2, t).font = Font(name=FONT, size=9, italic=True); k += 1
ws = wb.create_sheet('WinRate_vs_RR'); write_df(ws, WRR, fmts={'win_rate_gross': '0.0%', 'win_rate_net': '0.0%', 'theoretical_breakeven_wr': '0.0%', 'expectancy_R_gross': '0.000', 'expectancy_R_net': '0.000', 'trades': '#,##0'}, freeze=False)
ws.cell(len(WRR) + 3, 1, 'All strategies and parameter sets pooled, full sample, unselected. RR target = take-profit as a multiple of the stop. Hit rate falls as RR rises; net expectancy is negative everywhere once costs are paid.').font = Font(name=FONT, size=9, italic=True)
write_df(wb.create_sheet('Current_Params'), CUR, fmts={'win_rate_24m': '0.0%', 'pf_24m': '0.00', 'sharpe_24m': '0.00'}, widths={'strategy': 32, 'current_params': 46})
ws = wb['Current_Params']; ws.cell(len(CUR) + 3, 2, 'Fitted on 2024-06-16 to 2026-06-15 with the same selection rule. In-sample only: there is no out-of-sample period after this. Do not trade these without forward-testing.').font = Font(name=FONT, size=9, italic=True)
ws = wb.create_sheet('Portfolios'); write_df(ws, PORT, fmts={'sharpe': '0.00', 'total_ret_pct': '0.0', 'ann_ret_pct': '0.0', 'max_dd_pct': '0.0'}, widths={'portfolio': 66}, freeze=False)
ws.cell(len(PORT) + 3, 1, 'OOS period 2020-01-01 to 2026-06-15. Strategy returns are % of notional per strategy slot, equal weight, non-compounded; buy & hold is fully invested BTC (arithmetic daily sum).').font = Font(name=FONT, size=9, italic=True)

wb.move_sheet('Leaderboard', offset=0)
wb.save(f'{OUT}/BTC_15m_50_strategies_backtest.xlsx')

# ------------------------------------------------------------------ code bundle
z = zipfile.ZipFile(f'{OUT}/backtest_code.zip', 'w', zipfile.ZIP_DEFLATED)
for f in ['core.py', 'strategies.py', 'run_all.py', 'analyze.py', 'build_outputs.py']: z.write(f)
z.writestr('README.txt', 'Run order: python run_all.py ; python analyze.py ; python build_outputs.py\nEdit CSV path in core.py. Needs numpy pandas scipy numba matplotlib openpyxl.\n')
z.close()
print(len(TL), 'oos trades logged'); print(WRR.round(3).to_string())
