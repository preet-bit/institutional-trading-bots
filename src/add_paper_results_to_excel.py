import os, pickle, sys
import numpy as np, pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from core import run_bt, to_sig
from eurusd_pipeline import load_eurusd, build_windows, dayoff, pick, trades_of, sel, metrics, daily, sharpe, maxdd, START, END, FX_COST_RT, WINS, d0o, d1o, OOS_YEARS
from eurusd_papers_pipeline import PAPER_STRATS, SL_GRID, RR_GRID, MAX_HOLD, X

def main():
    print("Loading EURUSD data and running Paper Strategies...")
    df = load_eurusd()
    x = X(df)
    o, h, l, c = [df[k].values.astype(np.float64) for k in ('o', 'h', 'l', 'c')]
    a14 = x.atr(14).values.astype(np.float64)
    T = pd.to_datetime(df.t.values)
    DAY = ((T.floor('D') - START).days).values

    res = {}
    for st in PAPER_STRATS:
        for pi, p in enumerate(st['grid']):
            lg, sh = st['fn'](x, *p)
            sig = to_sig(lg, sh); sig[:150] = 0
            for sl in SL_GRID:
                for rr in RR_GRID:
                    res[(st['id'], pi, sl, rr)] = run_bt(o, h, l, c, sig, a14, sl, rr, MAX_HOLD)

    # Walk-forward across 28 windows
    WF_ROWS = []; OOS = {}
    for st in PAPER_STRATS:
        sid = st['id']; nets = []; rks = []; exds = []; ents = []; wid = []; grosses = []
        for wi, (a, b, c_win) in enumerate(WINS):
            da, db, dc = dayoff(a), dayoff(b), dayoff(c_win)
            best, breadth = pick(res, DAY, st, da, db, min_trades=40)
            row = dict(strategy_id=sid, strategy=st['name'], family=st['family'], window=wi + 1,
                       train=f"{a:%Y-%m-%d}..{b - pd.Timedelta(days=1):%Y-%m-%d}",
                       test=f"{b:%Y-%m-%d}..{c_win - pd.Timedelta(days=1):%Y-%m-%d}",
                       is_breadth_profitable=breadth)
            if best is None:
                row.update(selected='NO TRADE (nothing profitable in-sample)')
                WF_ROWS.append(row); continue
            (pi, sl, rr), im = best
            tr = trades_of(res, DAY, (sid, pi, sl, rr))
            m = sel(tr, db, dc)
            om = metrics(tr['net'][m], tr['risk'][m], tr['ex_day'][m], db, dc)
            row.update(selected=f"{st['grid'][pi]} SL={sl}xATR RR={rr}", sig_params=str(st['grid'][pi]),
                       param_idx=pi, sl_atr=sl, rr=rr, is_trades=im['trades'], is_win_rate=im['win_rate'],
                       is_sharpe=im['sharpe'], is_pf=im['profit_factor'], is_ret_pct=im['total_ret_pct'],
                       oos_trades=om['trades'], oos_win_rate=om['win_rate'], oos_pf=om['profit_factor'],
                       oos_ret_pct=om['total_ret_pct'], oos_exp_R=om['expectancy_R'], oos_sharpe=om['sharpe'])
            WF_ROWS.append(row)
            nets.append(tr['net'][m]); rks.append(tr['risk'][m]); exds.append(tr['ex_day'][m])
            ents.append(tr['ent_day'][m]); wid.append(np.full(m.sum(), wi))
            grosses.append(tr['gross'][m])
        cat = lambda L: np.concatenate(L) if L else np.array([])
        OOS[sid] = dict(net=cat(nets), gross=cat(grosses), risk=cat(rks),
                        exd=cat(exds).astype(int) if exds else np.array([], int),
                        entd=cat(ents), win=cat(wid))

    WF = pd.DataFrame(WF_ROWS)

    # Strategy Evaluation & Compounding Sizing
    rows = []
    yr_rows = []
    ydays = {y: (dayoff(pd.Timestamp(f'{y}-01-01')), dayoff(min(pd.Timestamp(f'{y + 1}-01-01'), END))) for y in range(2013, 2027)}

    for st in PAPER_STRATS:
        sid = st['id']; o_tr = OOS[sid]
        mt = metrics(o_tr['net'], o_tr['risk'], o_tr['exd'], d0o, d1o)
        dly = daily(o_tr['net'], o_tr['exd'], d0o, d1o) if len(o_tr['net']) else np.zeros(d1o - d0o)

        net = o_tr['net']; rk = o_tr['risk']
        if len(net) > 0 and len(rk) > 0:
            R = net / rk
            # 1.5% fixed risk compounded
            eq_15 = 100.0
            for r_val in R: eq_15 *= (1.0 + 0.015 * r_val)
            comp_15_ret = eq_15 - 100.0

            # 2.0% fixed risk compounded
            eq_20 = 100.0
            for r_val in R: eq_20 *= (1.0 + 0.020 * r_val)
            comp_20_ret = eq_20 - 100.0

            # 3x Leverage
            lev3_ret = dly.sum() * 3.0 * 100.0
        else:
            comp_15_ret = 0.0; comp_20_ret = 0.0; lev3_ret = 0.0

        # Window consistency
        act = WF[(WF.strategy_id == sid) & (WF.oos_trades.fillna(0) > 0)]
        traded_win = len(act)
        prof_win = (act.oos_ret_pct > 0).sum()
        pct_prof = prof_win / traded_win if traded_win else 0.0

        mt.update(
            strategy_id=sid, strategy=st['name'], family=st['family'],
            windows_traded=traded_win, windows_profitable=prof_win,
            pct_windows_profitable=pct_prof,
            unleveraged_1x_ret_pct=mt['total_ret_pct'],
            leverage_3x_ret_pct=lev3_ret,
            compounded_1_5pct_risk_pct=comp_15_ret,
            compounded_2_0pct_risk_pct=comp_20_ret,
            verdict='FAIL' if mt['sharpe'] < 0.50 else 'WATCH'
        )
        rows.append(mt)

        # Yearly returns
        yr_entry = dict(strategy_id=sid, strategy=st['name'])
        for y, (da_y, db_y) in ydays.items():
            m_y = (o_tr['exd'] >= da_y) & (o_tr['exd'] < db_y) if len(o_tr['net']) else np.array([], bool)
            yr_entry[str(y)] = o_tr['net'][m_y].sum() * 100.0 if len(o_tr['net']) else 0.0
        yr_rows.append(yr_entry)

    EV = pd.DataFrame(rows)
    YR = pd.DataFrame(yr_rows)

    print("\n========================= COMPLETE PAPER STRATEGY RESULTS =========================")
    cols_print = ['strategy_id', 'strategy', 'trades', 'win_rate', 'payoff', 'profit_factor',
                  'sharpe', 'max_dd_pct', 'pct_windows_profitable', 'unleveraged_1x_ret_pct',
                  'leverage_3x_ret_pct', 'compounded_1_5pct_risk_pct', 'verdict']
    print(EV[cols_print].to_string(index=False))

    # Append sheets to existing workbook
    xlsx_path = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'EURUSD_4h_50_strategies_backtest.xlsx'))
    wb = openpyxl.load_workbook(xlsx_path)

    FONT = 'Arial'
    hdr_fill = PatternFill('solid', fgColor='1F3864')
    hdr_font = Font(name=FONT, bold=True, color='FFFFFF', size=10)
    base = Font(name=FONT, size=10)
    thin = Side(style='thin', color='D9D9D9'); border = Border(bottom=thin)

    def write_sheet(ws, d, fmts=None):
        fmts = fmts or {}
        for j, c in enumerate(d.columns, 1):
            cell = ws.cell(1, j, str(c))
            cell.font = hdr_font; cell.fill = hdr_fill
            cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        ws.row_dimensions[1].height = 42
        for i, row in enumerate(d.itertuples(index=False), 2):
            for j, v in enumerate(row, 1):
                if isinstance(v, (np.floating, float)):
                    v = None if (np.isnan(v) or np.isinf(v)) else float(v)
                elif isinstance(v, np.integer): v = int(v)
                elif isinstance(v, (np.bool_, bool)): v = 'Y' if v else 'N'
                cell = ws.cell(i, j, v); cell.font = base; cell.border = border
                col_name = d.columns[j - 1]
                if col_name in fmts and v is not None:
                    cell.number_format = fmts[col_name]
        for j, c in enumerate(d.columns, 1):
            mx = max([len(str(c)) * .6] + [len(str(x)) for x in d[c].head(60)])
            ws.column_dimensions[get_column_letter(j)].width = min(max(9, mx + 2), 44)
        ws.freeze_panes = 'C2'
        return ws

    fm_std = {
        'win_rate': '0.0%', 'payoff': '0.00', 'profit_factor': '0.00', 'sharpe': '0.00',
        'sortino': '0.00', 'total_ret_pct': '0.0', 'ann_ret_pct': '0.0', 'max_dd_pct': '0.0',
        'pct_windows_profitable': '0.0%', 'unleveraged_1x_ret_pct': '0.0%', 'leverage_3x_ret_pct': '0.0%',
        'compounded_1_5pct_risk_pct': '0.0%', 'compounded_2_0pct_risk_pct': '0.0%'
    }

    # Remove old paper sheets if present
    for sname in ['Paper_Strategies_Leaderboard', 'Paper_Walkforward_Windows', 'Paper_Yearly_Returns']:
        if sname in wb.sheetnames:
            del wb[sname]

    ws_plb = wb.create_sheet('Paper_Strategies_Leaderboard')
    write_sheet(ws_plb, EV, fmts=fm_std)

    ws_pwf = wb.create_sheet('Paper_Walkforward_Windows')
    wfc = ['strategy_id', 'strategy', 'family', 'window', 'train', 'test', 'selected',
           'is_breadth_profitable', 'is_trades', 'is_win_rate', 'is_sharpe', 'is_pf', 'is_ret_pct',
           'oos_trades', 'oos_win_rate', 'oos_pf', 'oos_ret_pct', 'oos_exp_R', 'oos_sharpe']
    write_sheet(ws_pwf, WF[wfc], fmts={
        'is_breadth_profitable': '0.0%', 'is_win_rate': '0.0%', 'oos_win_rate': '0.0%',
        'is_sharpe': '0.00', 'is_pf': '0.00', 'is_ret_pct': '0.0', 'oos_pf': '0.00',
        'oos_ret_pct': '0.0', 'oos_exp_R': '0.000', 'oos_sharpe': '0.00'
    })

    ws_pyr = wb.create_sheet('Paper_Yearly_Returns')
    write_sheet(ws_pyr, YR, fmts={str(y): '0.0;[Red]-0.0' for y in range(2013, 2027)})

    wb.save(xlsx_path)
    print(f"Successfully saved updated Excel workbook with paper strategy sheets to: {xlsx_path}")
    out_xlsx = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'outputs', 'EURUSD_4h_50_strategies_backtest.xlsx'))
    wb.save(out_xlsx)
    print(f"Copied to outputs: {out_xlsx}")

if __name__ == '__main__':
    main()
